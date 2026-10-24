"""Prepare short quotation excerpts and editable Markdown copies for review.

No canonical writes, imports, provider calls or automatic translation decisions.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import html
import importlib.util
import json
from pathlib import Path
import re
import subprocess

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill


REPO_ROOT = Path(__file__).resolve().parents[1]
LANGUAGES = ("de", "fr", "it")


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def quotation_text(raw: str) -> str:
    value = html.unescape(raw)
    if value.startswith(r"\glqq"):
        value = re.sub(r"^\\glqq(?:\{\})?\s*", "", value)
        value = re.sub(r"\s*\\grqq(?:\{\})?$", "", value)
    else:
        value = value[1:-1]
    return value.strip()


def excerpt(value: str) -> str:
    visible = re.sub(r"\s+", " ", value).strip()
    return visible if len(visible) <= 15 else visible[:15] + "…"


def load_inventory(translation_tool_root: Path):
    path = translation_tool_root / "question_pipeline" / "source_format_contract.py"
    if not path.is_file():
        raise ValueError(f"Shared quotation extractor not found: {path}")
    spec = importlib.util.spec_from_file_location("quotation_review_source_format", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.formatting_inventory, path


def prepare(*, canonical_root: Path, candidates_manifest: Path, output_dir: Path,
            translation_tool_root: Path) -> dict:
    canonical_root = canonical_root.resolve()
    candidates_manifest = candidates_manifest.resolve()
    output_dir = output_dir.resolve()
    if output_dir.exists() or output_dir.is_relative_to(canonical_root):
        raise ValueError("Output directory must be new and outside canonical.")
    candidate_data = json.loads(candidates_manifest.read_text(encoding="utf-8"))
    screen = candidate_data["post_change_quote_pair_count_screen"]
    candidates = screen["candidates"]
    if len(candidates) != screen["candidate_count"]:
        raise ValueError("Candidate count disagrees with its manifest.")
    inventory, extractor_path = load_inventory(translation_tool_root)
    objects = []
    seen = set()
    for candidate in candidates:
        folder = (canonical_root / candidate["object"]).resolve()
        if not folder.is_relative_to(canonical_root):
            raise ValueError(f"Candidate outside canonical: {folder}")
        meta_path = folder / "object.meta.json"
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
        object_id = meta["id"]
        if folder.name != object_id or object_id in seen:
            raise ValueError(f"Duplicate or inconsistent candidate ID: {object_id}")
        seen.add(object_id)
        files = {}
        for language in LANGUAGES:
            path = folder / f"body.{language}.md"
            original = path.read_bytes()
            text = original.decode("utf-8")
            quotes = [
                {**span, "quoted_text": quotation_text(span["text"])}
                for span in inventory(text)["quotation_pairs"]
            ]
            files[language] = {
                "canonical_path": str(path.relative_to(canonical_root)),
                "review_path": f"{object_id}/{path.name}",
                "sha256": digest(original),
                "original_text": text,
                "quotation_pairs": quotes,
            }
        objects.append({
            "id": object_id, "object_path": candidate["object"],
            "source_md": meta.get("source", {}).get("path", "body.de.md"),
            "metadata_sha256": digest(meta_path.read_bytes()),
            "historical_screen_counts": candidate["paired_quote_count"], "files": files,
        })

    # Recheck the complete input snapshot before creating editable artifacts.
    for obj in objects:
        for item in obj["files"].values():
            assert digest((canonical_root / item["canonical_path"]).read_bytes()) == item["sha256"]
    output_dir.mkdir(parents=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Guillemets"
    sheet.append(["ID", "Fichier MD allemand", "DE", "FR", "IT"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7C2620")
    sheet.freeze_panes = "C2"
    for column, width in {"A": 26, "B": 38, "C": 25, "D": 25, "E": 25}.items():
        sheet.column_dimensions[column].width = width
    for object_index, obj in enumerate(objects):
        for item in obj["files"].values():
            target = output_dir / item["review_path"]
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(item["original_text"].encode("utf-8"))
        count = max(1, *(len(item["quotation_pairs"]) for item in obj["files"].values()))
        obj["excel_first_row"] = sheet.max_row + 1
        for occurrence in range(count):
            row = sheet.max_row + 1
            sheet.cell(row, 1, obj["id"])
            sheet.cell(row, 2, Path(obj["source_md"]).name)
            sheet.cell(row, 2).hyperlink = obj["files"]["de"]["review_path"]
            sheet.cell(row, 2).style = "Hyperlink"
            counts = {lang: len(obj["files"][lang]["quotation_pairs"]) for lang in LANGUAGES}
            sheet.cell(row, 1).comment = Comment(
                f"Occurrence {occurrence + 1} dans chaque langue.\n"
                f"Paires reconnues : {counts}\n"
                "Les colonnes suivent l'ordre d'apparition, sans garantie d'équivalence sémantique.",
                "Revue guillemets")
            sheet.cell(row, 2).comment = Comment(
                f"Source allemande : {obj['source_md']}\n"
                f"Copie à modifier : {obj['files']['de']['review_path']}", "Revue guillemets")
            for column, language in enumerate(LANGUAGES, 3):
                item = obj["files"][language]
                pairs = item["quotation_pairs"]
                cell = sheet.cell(row, column)
                cell.hyperlink = item["review_path"]
                if occurrence < len(pairs):
                    pair = pairs[occurrence]
                    cell.value = excerpt(pair["quoted_text"])
                    cell.data_type = "s"
                    cell.comment = Comment(
                        f"{item['review_path']} — ligne {pair['line']}\n\n"
                        f"Texte complet entre guillemets :\n{pair['quoted_text']}\n\n"
                        f"Écriture d'origine :\n{pair['text']}", "Revue guillemets")
                else:
                    cell.value = "—"
                    cell.comment = Comment(
                        "Aucune paire reconnue à ce rang. Cela ne prouve ni l'absence du passage traduit, "
                        "ni une faute de traduction. Vérifier le fichier complet.", "Revue guillemets")
                cell.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
            for cell in sheet[row]:
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                if object_index % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F5F7")
            sheet.row_dimensions[row].height = 24
        obj["excel_last_row"] = sheet.max_row
    sheet.auto_filter.ref = sheet.dimensions
    instructions = workbook.create_sheet("Mode d'emploi")
    instructions.column_dimensions["A"].width = 115
    notes = [
        "Modifier les copies Markdown dans les sous-dossiers, pas les extraits dans l'Excel.",
        "Chaque cellule DE/FR/IT montre le texte entre guillemets, limité à 15 caractères puis ….",
        "Le commentaire de cellule donne l'expression complète, son écriture et le numéro de ligne.",
        "Les liens relatifs ouvrent les copies Markdown ; conserver le classeur dans ce dossier.",
        "Plusieurs paires donnent plusieurs lignes par objet, par ordre d'apparition dans chaque langue.",
        "Le rang n'est pas un alignement sémantique : une paire omise peut décaler les expressions suivantes.",
        "— signifie seulement qu'aucune paire n'est reconnue à ce rang, pas que le texte est absent.",
        "Le même extracteur déterministe que celui des traductions est utilisé ; il exclut code, attributs et commentaires.",
        "Les 35 candidats historiques sont conservés même si leur inventaire recalculé diffère du premier comptage.",
        "Aucune modification des copies n'est réintégrée automatiquement dans le canonical.",
        "Le manifeste conserve les originaux et leurs empreintes pour contrôler les conflits lors de la réintégration.",
    ]
    for note in notes:
        instructions.append([note])
        instructions.cell(instructions.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        instructions.row_dimensions[instructions.max_row].height = 34
    workbook_path = output_dir / "quotation-review.xlsx"
    workbook.save(workbook_path)
    git = subprocess.run(["git", "-C", str(canonical_root), "rev-parse", "HEAD"],
                         capture_output=True, text=True, check=False)
    manifest = {
        "schema_version": 1, "kind": "editable_quotation_manual_review",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "canonical_root": str(canonical_root), "canonical_git_commit": git.stdout.strip(),
        "baseline": "current canonical working tree, including pending accepted changes",
        "candidates_manifest": str(candidates_manifest),
        "candidates_manifest_sha256": digest(candidates_manifest.read_bytes()),
        "extractor": str(extractor_path), "extractor_sha256": digest(extractor_path.read_bytes()),
        "preparation_tool_sha256": digest(Path(__file__).read_bytes()),
        "workbook": workbook_path.name, "initial_workbook_sha256": digest(workbook_path.read_bytes()),
        "object_count": len(objects), "quotation_rows": sheet.max_row - 1,
        "snippet_limit": 15, "row_alignment": "independent occurrence order, not semantic matching",
        "reintegration": "Only after explicit instruction; compare each changed copy with original_text and require the current canonical hash to match sha256 before applying. Conflicts need explicit resolution. Preserve all human edits and all review artifacts.",
        "objects": objects,
    }
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    workspace = {"folders": [{"name": "Revue des guillemets — copies Markdown", "path": "."}]}
    (output_dir / "quotation-review.code-workspace").write_text(json.dumps(workspace, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output_dir / "README.md").write_text(
        "# Revue manuelle des guillemets\n\n"
        f"{len(objects)} objets, trois copies Markdown par objet. Ouvrir `quotation-review.code-workspace` dans VS Code.\n\n"
        + "\n".join("- " + note for note in notes)
        + "\n\nLe nom affiché dans l'Excel est le nom du fichier allemand d'origine ; son lien mène à la copie `body.de.md`. "
          "Conserver les noms des dossiers et fichiers et ne pas éditer `manifest.json`. Les trois langues peuvent être corrigées. "
          "Lors de la reprise, les changements seront d'abord présentés et les conflits contrôlés avant toute écriture canonical.\n",
        encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--canonical-root", type=Path, default=REPO_ROOT / "canonical")
    parser.add_argument("--candidates-manifest", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--translation-tool-root", type=Path, default=REPO_ROOT.parent / "50ohm-ai-translation-tool")
    args = parser.parse_args()
    result = prepare(canonical_root=args.canonical_root, candidates_manifest=args.candidates_manifest,
                     output_dir=args.output_dir, translation_tool_root=args.translation_tool_root)
    print(json.dumps({"objects": result["object_count"], "quotation_rows": result["quotation_rows"],
                      "output_dir": str(args.output_dir.resolve())}))


if __name__ == "__main__":
    main()
