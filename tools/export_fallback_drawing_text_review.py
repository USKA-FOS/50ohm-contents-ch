#!/usr/bin/env python
"""Export one row per visible-text/drawing tuple using German SVG fallback."""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import extract_drawing_tex_translation_candidates as extractor


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REVIEW_CSV = (
    REPO_ROOT / "work" / "drawing_text_audit" / "drawing_tex_translation_candidates_2.csv"
)
DEFAULT_OUTPUT = (
    REPO_ROOT / "work" / "drawing_text_audit" / "fallback_drawing_german_text_review.xlsx"
)
HEADERS = (
    "text",
    "drawing_number",
    "canonical_id",
    "canonical_reference",
    "source_tex",
    "candidate_origin",
    "category",
    "fr_svg_status",
    "it_svg_status",
    "decision",
    "reviewer",
    "comment",
)


def read_previous_statuses(path: Path) -> dict[tuple[str, str], set[str]]:
    statuses: dict[tuple[str, str], set[str]] = defaultdict(set)
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            statuses[(row["canonical_reference"], row["translatable_text"])].add(
                row["to_be_translated"].strip().lower()
            )
    return statuses


def build_rows(review_csv: Path) -> list[list[str]]:
    previous = read_previous_statuses(review_csv)
    localized = {
        language: {
            str(path.parent.relative_to(REPO_ROOT))
            for path in extractor.CANONICAL_DRAWINGS.glob(f"*/*.{language}.svg")
        }
        for language in ("fr", "it")
    }
    rows: list[list[str]] = []
    seen: set[tuple[str, str]] = set()
    for candidate in extractor.build_candidates():
        reference = candidate.canonical_reference
        if reference in localized["fr"] and reference in localized["it"]:
            continue
        key = (candidate.translatable_text, reference)
        if key in seen:
            continue
        seen.add(key)
        statuses = previous.get((reference, candidate.translatable_text), set())
        if "true" in statuses:
            origin = "previously_marked_to_translate"
        elif statuses and statuses <= {"false"}:
            continue
        else:
            origin = "newly_detected"
        object_dir = REPO_ROOT / reference
        source_tex = next(object_dir.glob("*.de.tex")).name
        rows.append(
            [
                candidate.translatable_text,
                candidate.figure_number,
                Path(reference).name,
                reference,
                source_tex,
                origin,
                candidate.category,
                "explicit" if reference in localized["fr"] else "fallback_de",
                "explicit" if reference in localized["it"] else "fallback_de",
                "",
                "",
                "",
            ]
        )

    def sort_key(row: list[str]) -> tuple[str, int, int | str]:
        number = row[1]
        number_key: tuple[int, int | str] = (
            (0, int(number)) if number.isdigit() else (1, number.casefold())
        )
        return (row[0].casefold(), *number_key)

    return sorted(rows, key=sort_key)


def write_workbook(output: Path, review_csv: Path) -> int:
    rows = build_rows(review_csv)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "fallback_text_review"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = (42, 16, 20, 40, 18, 32, 24, 16, 16, 18, 16, 40)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_view.showGridLines = False
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--review-csv", type=Path, default=DEFAULT_REVIEW_CSV)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    row_count = write_workbook(args.output, args.review_csv)
    print(f"output={args.output.resolve()}")
    print(f"tuple_count={row_count}")


if __name__ == "__main__":
    main()
