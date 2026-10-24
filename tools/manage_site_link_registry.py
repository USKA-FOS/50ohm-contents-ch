#!/usr/bin/env python
"""Inventory multilingual site links and check their reachability.

The German href is the stable registry key. Fragment-only hrefs are scoped to
their source page because ``#section`` is meaningful only within that page.
"""

from __future__ import annotations

import argparse
import email.utils
import hashlib
import json
import re
import urllib.error
import urllib.parse
import urllib.request
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import UTC, datetime
from difflib import SequenceMatcher
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SITE_ROOT = ROOT.parent / "50ohm-site-releases"
DEFAULT_WORKBOOK = ROOT / "review" / "link_localization" / "site_link_registry.xlsx"
DEFAULT_EXTERNAL_REPORT = ROOT / "review" / "link_localization" / "external_site_link_report.xlsx"
LANGUAGES = ("de", "fr", "it")
STATUSES = {"not found", "found", "verified"}
LINK_COLUMNS = {lang: f"Link {lang.upper()}" for lang in LANGUAGES}
TEXT_COLUMNS = {lang: f"Text {lang.upper()}" for lang in LANGUAGES}
STATUS_COLUMNS = {lang: f"Status {lang.upper()}" for lang in LANGUAGES}


@dataclass
class Anchor:
    href: str
    text_parts: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return " ".join(" ".join(self.text_parts).split())


class AnchorParser(HTMLParser):
    """Extract every anchor and its readable text from generated HTML."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.anchors: list[Anchor] = []
        self.current: Anchor | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() == "a":
            href = dict(attrs).get("href")
            if href is not None:
                self.current = Anchor(href=href)
                self.anchors.append(self.current)
        elif tag.lower() == "img" and self.current:
            alt = dict(attrs).get("alt")
            if alt:
                self.current.text_parts.append(alt)

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_data(self, data: str) -> None:
        if self.current:
            self.current.text_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "a":
            self.current = None


@dataclass
class Occurrence:
    link_hash: str
    href_de: str
    page: str
    ordinal_de: int
    text_de: str
    href_fr: str = ""
    text_fr: str = ""
    match_fr: str = "unmatched"
    href_it: str = ""
    text_it: str = ""
    match_it: str = "unmatched"


def link_hash(href: str, page: str) -> str:
    """Hash a German href, scoping a bare fragment to its source page."""
    key = f"{page}\0{href}" if href.startswith("#") else href
    return hashlib.sha256(key.encode("utf-8")).hexdigest()


def link_type(href: str, site_origin: str = "https://50ohm.uska.ch") -> str:
    """Classify a site-relative or same-site URL as internal."""
    parsed = urllib.parse.urlsplit(href)
    if parsed.scheme in {"mailto", "tel"}:
        return "external"
    if parsed.scheme in {"http", "https"}:
        return "internal" if parsed.netloc.casefold() == urllib.parse.urlsplit(site_origin).netloc.casefold() else "external"
    if parsed.netloc:
        return "internal" if parsed.netloc.casefold() == urllib.parse.urlsplit(site_origin).netloc.casefold() else "external"
    return "internal"


def is_relative_internal_href(href: str) -> bool:
    parsed = urllib.parse.urlsplit(href)
    return not parsed.scheme and not parsed.netloc


def extract_page(path: Path) -> list[Anchor]:
    parser = AnchorParser()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    return parser.anchors


def align_anchors(source: list[Anchor], target: list[Anchor]) -> dict[int, int]:
    """Map source anchor positions to target positions without shifting gaps."""
    if len(source) == len(target):
        return dict(enumerate(range(len(target))))

    source_hrefs = [a.href for a in source]
    target_hrefs = [a.href for a in target]
    matcher = SequenceMatcher(a=source_hrefs, b=target_hrefs, autojunk=False)
    aligned: dict[int, int] = {}
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        source_count, target_count = i2 - i1, j2 - j1
        if tag == "equal":
            aligned.update((i1 + offset, j1 + offset) for offset in range(source_count))
        elif tag == "replace" and source_count == target_count:
            aligned.update((i1 + offset, j1 + offset) for offset in range(source_count))
    return aligned


def load_sites(site_root: Path) -> dict[str, dict[str, list[Anchor]]]:
    sites: dict[str, dict[str, list[Anchor]]] = {}
    for lang in LANGUAGES:
        language_root = site_root / lang
        if not language_root.is_dir():
            raise FileNotFoundError(f"Missing generated site directory: {language_root}")
        pages = {}
        # The deployed page set is in the language root. Nested HTML files are
        # bundled vendor demos/tests (for example Reveal.js examples), not site
        # content pages and not part of the link-review scope.
        for path in sorted(language_root.glob("*.html")):
            pages[path.relative_to(language_root).as_posix()] = extract_page(path)
        sites[lang] = pages
    return sites


def read_previous_registry(path: Path) -> dict[str, dict[str, object]]:
    if not path.exists():
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Links" not in workbook.sheetnames:
        raise ValueError(f"Workbook has no Links sheet: {path}")
    sheet = workbook["Links"]
    headers = {cell.value: cell.column for cell in sheet[1] if cell.value}
    if "German link SHA-256" not in headers:
        raise ValueError(f"Workbook has no German link SHA-256 key: {path}")
    previous = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        key = row[headers["German link SHA-256"] - 1]
        if not key:
            continue
        record = {}
        for lang in LANGUAGES:
            for column in (LINK_COLUMNS[lang], STATUS_COLUMNS[lang]):
                if column in headers:
                    record[column] = row[headers[column] - 1]
        previous[str(key)] = record
    workbook.close()
    return previous


def unique_lines(values: list[str]) -> str:
    return "\n".join(dict.fromkeys(value for value in values if value))


def export_registry(site_root: Path, output: Path, previous_path: Path | None) -> tuple[int, int]:
    sites = load_sites(site_root)
    previous = read_previous_registry(previous_path) if previous_path and previous_path.exists() else {}
    occurrence_rows: list[Occurrence] = []
    lang_alignment_counts = {"fr": 0, "it": 0}

    for page, source_anchors in sites["de"].items():
        target_anchors = {lang: sites[lang].get(page, []) for lang in ("fr", "it")}
        aligned = {lang: align_anchors(source_anchors, target_anchors[lang]) for lang in ("fr", "it")}
        lang_alignment_counts.update(
            {
                lang: lang_alignment_counts[lang] + len(aligned[lang])
                for lang in ("fr", "it")
            }
        )
        for ordinal, anchor in enumerate(source_anchors):
            row = Occurrence(
                link_hash(anchor.href, page), anchor.href, page, ordinal,
                anchor.text,
            )
            for lang in ("fr", "it"):
                target_index = aligned[lang].get(ordinal)
                if target_index is not None:
                    target = target_anchors[lang][target_index]
                    setattr(row, f"href_{lang}", target.href)
                    setattr(row, f"text_{lang}", target.text)
                    setattr(row, f"match_{lang}", "matched")
            occurrence_rows.append(row)

    grouped: dict[str, list[Occurrence]] = defaultdict(list)
    for row in occurrence_rows:
        grouped[row.link_hash].append(row)

    try:
        manifest = json.loads((site_root / "release-manifest.json").read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        manifest = {}
    release_id = manifest.get("release_id", "unknown")
    created_at = datetime.now(UTC).isoformat(timespec="seconds")

    workbook = Workbook()
    links_sheet = workbook.active
    links_sheet.title = "Links"
    headers = [
        "German link SHA-256",
        "German link key",
        "Link type",
        "Text DE",
        "Text FR",
        "Text IT",
        "Link DE",
        "Link FR",
        "Link IT",
        "Suggested link FR",
        "Suggested link IT",
        "Status DE",
        "Status FR",
        "Status IT",
        "DE occurrences",
    ]
    links_sheet.append(headers)
    for key, rows in sorted(grouped.items(), key=lambda item: (item[1][0].href_de, item[1][0].page)):
        href_de = rows[0].href_de
        texts = {"de": [r.text_de for r in rows], "fr": [r.text_fr for r in rows], "it": [r.text_it for r in rows]}
        hrefs = {
            "de": [r.href_de for r in rows],
            "fr": [r.href_fr for r in rows],
            "it": [r.href_it for r in rows],
        }
        prior = previous.get(key, {})
        href_cells = {lang: unique_lines(hrefs[lang]) for lang in LANGUAGES}
        suggested = {"fr": "", "it": ""}
        if is_relative_internal_href(href_de):
            for lang in ("fr", "it"):
                if not href_cells[lang]:
                    suggested[lang] = href_de
        status_cells = {lang: "" for lang in LANGUAGES}
        for lang in LANGUAGES:
            if prior.get(STATUS_COLUMNS[lang]) == "verified":
                old_href = str(prior.get(LINK_COLUMNS[lang]) or "")
                if lang == "de" or not old_href or old_href == href_cells[lang]:
                    status_cells[lang] = "verified"
                elif not href_cells[lang]:
                    suggested[lang] = old_href
        links_sheet.append(
            [
                key,
                href_de,
                link_type(href_de),
                unique_lines(texts["de"]),
                unique_lines(texts["fr"]),
                unique_lines(texts["it"]),
                href_cells["de"],
                href_cells["fr"],
                href_cells["it"],
                suggested["fr"],
                suggested["it"],
                status_cells["de"],
                status_cells["fr"],
                status_cells["it"],
                len(rows),
            ]
        )

    occurrence_sheet = workbook.create_sheet("Occurrences")
    occurrence_headers = [
        "German link SHA-256", "German page", "German anchor order",
        "Text DE", "Link DE", "Text FR", "Observed link FR", "FR match",
        "Text IT", "Observed link IT", "IT match",
    ]
    occurrence_sheet.append(occurrence_headers)
    for row in occurrence_rows:
        occurrence_sheet.append(
            [
                row.link_hash, row.page, row.ordinal_de,
                row.text_de, row.href_de,
                row.text_fr, row.href_fr, row.match_fr,
                row.text_it, row.href_it, row.match_it,
            ]
        )

    manifest_sheet = workbook.create_sheet("Manifest")
    manifest_sheet.append(["Field", "Value"])
    manifest_values = [
        ("Release ID", release_id),
        ("Created at (UTC)", created_at),
        ("Release manifest", "release-manifest.json"),
        ("German source commit", manifest.get("sources", {}).get("contents", {}).get("commit", "")),
        ("Generator commit", manifest.get("sources", {}).get("generator", {}).get("commit", "")),
        ("Matching scope", "All <a href> anchors in root-level generated page HTML; nested third-party assets are excluded. Target anchors align by page and order."),
        ("Hash key rule", "SHA-256 of exact German href; bare #fragment hrefs also include the German source page."),
        ("German HTML pages", len(sites["de"])),
        ("French HTML pages", len(sites["fr"])),
        ("Italian HTML pages", len(sites["it"])),
        ("Unique German link keys", len(grouped)),
        ("German anchor occurrences", len(occurrence_rows)),
        ("FR aligned occurrences", lang_alignment_counts["fr"]),
        ("IT aligned occurrences", lang_alignment_counts["it"]),
    ]
    for row in manifest_values:
        manifest_sheet.append(row)

    style_sheet(links_sheet, widths=[66, 50, 14, 40, 40, 40, 52, 52, 52, 52, 52, 16, 16, 16, 16])
    style_sheet(occurrence_sheet, widths=[66, 42, 16, 36, 52, 36, 52, 14, 36, 52, 14])
    style_sheet(manifest_sheet, widths=[32, 115])
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return len(grouped), len(occurrence_rows)


def style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A75")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def split_cell_urls(value: object) -> list[str]:
    if value is None:
        return []
    return list(dict.fromkeys(line.strip() for line in str(value).splitlines() if line.strip()))


def is_reveal_route_fragment(fragment: str) -> bool:
    return bool(re.fullmatch(r"\d+(?:/\d+)?", fragment))


def ascii_http_url(url: str) -> str:
    """Percent-encode Unicode URL components and IDNA-encode its host."""
    parts = urllib.parse.urlsplit(url)
    hostname = parts.hostname or ""
    if ":" in hostname:
        ascii_host = f"[{hostname}]"
    else:
        ascii_host = hostname.encode("idna").decode("ascii")
    if parts.port is not None:
        ascii_host = f"{ascii_host}:{parts.port}"
    if parts.username is not None:
        userinfo = urllib.parse.quote(urllib.parse.unquote(parts.username), safe="-._~")
        if parts.password is not None:
            password = urllib.parse.quote(urllib.parse.unquote(parts.password), safe="-._~")
            userinfo += f":{password}"
        ascii_host = f"{userinfo}@{ascii_host}"
    path = urllib.parse.quote(urllib.parse.unquote(parts.path), safe="/%:@!$&'()*+,;=-._~")
    query = urllib.parse.quote(urllib.parse.unquote(parts.query), safe="/?@!$&'()*+,;=:%-._~")
    return urllib.parse.urlunsplit((parts.scheme, ascii_host, path, query, ""))


def read_fragment_ids(path: Path) -> set[str]:
    parser = IdParser()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    return parser.ids


class IdParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.ids: set[str] = set()

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        values = dict(attrs)
        for name in ("id", "name"):
            if values.get(name):
                self.ids.add(str(values[name]))


class LinkChecker:
    def __init__(self, site_root: Path, site_origin: str, timeout: float, retry_429: int = 1,
                 max_retry_wait: float = 30.0) -> None:
        self.site_root = site_root.resolve()
        self.site_origin = site_origin.rstrip("/")
        self.timeout = timeout
        self.retry_429 = max(0, retry_429)
        self.max_retry_wait = max(0.0, max_retry_wait)
        self.fragment_cache: dict[Path, set[str]] = {}
        self.reveal_page_cache: dict[Path, bool] = {}
        self.http_cache: dict[str, tuple[str, str, str]] = {}

    def check(self, href: str, lang: str, source_pages: list[str]) -> tuple[str, str, str]:
        if not href:
            return "not found", "", "Empty href"
        parsed = urllib.parse.urlsplit(href)
        if parsed.scheme in {"http", "https"}:
            if link_type(href, self.site_origin) == "internal":
                return self.check_internal_absolute(href, lang, source_pages)
            return self.check_http(urllib.parse.urlunsplit(parsed._replace(fragment="")))
        if parsed.netloc and link_type(href, self.site_origin) == "internal":
            return self.check_internal_absolute(urllib.parse.urlunsplit(("https", parsed.netloc, parsed.path, parsed.query, parsed.fragment)), lang, source_pages)
        if parsed.scheme in {"mailto", "tel"}:
            return "found", "", f"Syntactically valid {parsed.scheme}: URI; destination itself cannot be probed"
        if parsed.scheme:
            return f"unsupported scheme: {parsed.scheme}", "", f"Unsupported URI scheme: {parsed.scheme}"

        pages = source_pages or ["index.html"]
        results = [self.check_local(href, lang, page) for page in pages]
        if any(result[0] == "not found" for result in results):
            detail = "; ".join(dict.fromkeys(result[2] for result in results if result[0] == "not found"))
            return "not found", "", detail
        if results and all(result[0] == "found" for result in results):
            return "found", "", f"Resolved in {len(results)} source-page context(s)"
        detail = "; ".join(dict.fromkeys(result[2] for result in results if result[2]))
        return "local check inconclusive", "", detail or "Could not resolve local target"

    def check_internal_absolute(self, href: str, lang: str, source_pages: list[str]) -> tuple[str, str, str]:
        parsed = urllib.parse.urlsplit(href)
        origin_path = urllib.parse.urlsplit(self.site_origin).path.rstrip("/")
        path = parsed.path
        if origin_path and path.startswith(origin_path + "/"):
            path = path[len(origin_path):]
        # Deployed paths may include a prefix such as /isq/site before /fr/.
        parts = [part for part in path.split("/") if part]
        lang_index = next((index for index, part in enumerate(parts) if part in LANGUAGES), None)
        if lang_index is not None:
            local_path = "/" + "/".join(parts[lang_index:])
        else:
            local_path = path
        local_href = urllib.parse.urlunsplit(("", "", local_path or "/", parsed.query, parsed.fragment))
        results = [self.check_local(local_href, lang, page) for page in (source_pages or ["index.html"])]
        states = {result[0] for result in results}
        if "not found" in states:
            return "not found", "", "; ".join(dict.fromkeys(r[2] for r in results if r[0] == "not found"))
        if results and all(result[0] == "found" for result in results):
            return "found", "", "; ".join(dict.fromkeys(r[2] for r in results))
        return "local check inconclusive", "", "; ".join(dict.fromkeys(r[2] for r in results))

    def check_local(self, href: str, lang: str, source_page: str) -> tuple[str, str, str]:
        base = f"https://local.invalid/{lang}/{urllib.parse.quote(source_page, safe='/')}"
        resolved = urllib.parse.urlsplit(urllib.parse.urljoin(base, href))
        if resolved.netloc != "local.invalid":
            return self.check_http(urllib.parse.urlunsplit(resolved._replace(fragment="")))
        parts = [part for part in resolved.path.split("/") if part]
        target_lang = lang
        if parts and parts[0] in LANGUAGES:
            target_lang = parts.pop(0)
        relative = Path(*[urllib.parse.unquote(part) for part in parts]) if parts else Path("index.html")
        language_root = (self.site_root / target_lang).resolve()
        target = (language_root / relative).resolve()
        if language_root not in target.parents and target != language_root:
            return "not found", "", "Resolved path escapes the release site"
        if target.is_dir():
            target = target / "index.html"
        if not target.is_file():
            return "not found", "", f"Local target does not exist: {target}"
        if resolved.fragment:
            fragment = urllib.parse.unquote(resolved.fragment)
            if self.is_reveal_page(target) and is_reveal_route_fragment(fragment):
                return "found", "", f"Reveal.js slide route #{fragment} in {target.name}"
            if target not in self.fragment_cache:
                self.fragment_cache[target] = read_fragment_ids(target)
            if fragment not in self.fragment_cache[target]:
                return "not found", "", f"Fragment #{fragment} not found in {target.name}"
        return "found", "", f"Local target exists: {target.relative_to(self.site_root) if self.site_root in target.parents else target}"

    def is_reveal_page(self, path: Path) -> bool:
        if path not in self.reveal_page_cache:
            source = path.read_text(encoding="utf-8", errors="replace")
            self.reveal_page_cache[path] = 'class="reveal' in source or "Reveal.initialize" in source
        return self.reveal_page_cache[path]

    def check_http(self, url: str) -> tuple[str, str, str]:
        if url in self.http_cache:
            return self.http_cache[url]
        try:
            request_url = ascii_http_url(url)
        except (UnicodeError, ValueError) as error:
            result = ("invalid URL", "", f"Invalid HTTP URL: {error}")
            self.http_cache[url] = result
            return result
        for attempt in range(self.retry_429 + 1):
            request = urllib.request.Request(
                request_url,
                headers={"User-Agent": "50ohm-link-registry/1.0 (+site link accessibility audit)"},
            )
            try:
                with urllib.request.urlopen(request, timeout=self.timeout) as response:
                    response.read(1)
                    code = str(response.status)
                    result = ("found", code, f"HTTP {code}; final URL {response.geturl()}")
                break
            except urllib.error.HTTPError as error:
                code = str(error.code)
                if error.code in {404, 410}:
                    result = ("not found", code, f"HTTP {code}: {error.reason}")
                elif error.code == 429 and attempt < self.retry_429:
                    retry_after = error.headers.get("Retry-After", "")
                    try:
                        delay = max(0.0, float(retry_after))
                    except ValueError:
                        try:
                            retry_at = email.utils.parsedate_to_datetime(retry_after)
                            delay = max(0.0, (retry_at - datetime.now(UTC)).total_seconds())
                        except (TypeError, ValueError, OverflowError):
                            delay = 15.0
                    if delay > self.max_retry_wait:
                        result = ("HTTP 429", code, f"HTTP {code}: Retry-After={retry_after!r} exceeds the configured wait limit")
                        break
                    time.sleep(delay)
                    continue
                else:
                    result = (f"HTTP {code}", code, f"HTTP {code}: {error.reason}")
                break
            except urllib.error.URLError as error:
                reason = error.reason
                status = "timeout" if isinstance(reason, TimeoutError) else "network error"
                result = (status, "", f"{status}: {reason}")
                break
            except (TimeoutError, OSError) as error:
                result = ("timeout" if isinstance(error, TimeoutError) else "network error", "", str(error))
                break
        self.http_cache[url] = result
        return result


def check_registry(workbook_path: Path, site_root: Path, site_origin: str, timeout: float, workers: int,
                   retry_429: int = 1, max_retry_wait: float = 30.0) -> dict[str, int]:
    workbook = load_workbook(workbook_path)
    if "Links" not in workbook.sheetnames or "Occurrences" not in workbook.sheetnames:
        raise ValueError("Expected Links and Occurrences sheets; run the scan command first")
    links_sheet = workbook["Links"]
    occurrences_sheet = workbook["Occurrences"]
    link_headers = {cell.value: cell.column for cell in links_sheet[1] if cell.value}
    occurrence_headers = {cell.value: cell.column for cell in occurrences_sheet[1] if cell.value}
    required_link_headers = {"German link SHA-256", *LINK_COLUMNS.values(), *STATUS_COLUMNS.values()}
    if not required_link_headers.issubset(link_headers):
        raise ValueError(f"Missing workbook columns: {sorted(required_link_headers - set(link_headers))}")

    pages_by_key: dict[str, dict[str, dict[str, list[str]]]] = defaultdict(
        lambda: {lang: defaultdict(list) for lang in LANGUAGES}
    )
    for row in occurrences_sheet.iter_rows(min_row=2, values_only=True):
        key = row[occurrence_headers["German link SHA-256"] - 1]
        page = row[occurrence_headers["German page"] - 1]
        if not key or not page:
            continue
        for lang, header in (("de", "Link DE"), ("fr", "Observed link FR"), ("it", "Observed link IT")):
            href = row[occurrence_headers[header] - 1]
            if href and str(page) not in pages_by_key[str(key)][lang][str(href)]:
                pages_by_key[str(key)][lang][str(href)].append(str(page))

    checker = LinkChecker(site_root, site_origin, timeout, retry_429, max_retry_wait)
    work: list[tuple[int, str, str, str, list[str]]] = []
    for row_index in range(2, links_sheet.max_row + 1):
        key = links_sheet.cell(row_index, link_headers["German link SHA-256"]).value
        if not key:
            continue
        for lang in LANGUAGES:
            hrefs = split_cell_urls(links_sheet.cell(row_index, link_headers[LINK_COLUMNS[lang]]).value)
            prior = links_sheet.cell(row_index, link_headers[STATUS_COLUMNS[lang]]).value or ""
            fallback_pages = list(
                dict.fromkeys(
                    page
                    for page_list in pages_by_key[str(key)]["de"].values()
                    for page in page_list
                )
            )
            source_href = links_sheet.cell(row_index, link_headers[LINK_COLUMNS["de"]]).value
            if not hrefs:
                missing_status = "missing link" if link_type(str(source_href or ""), site_origin) == "internal" else "missing external link"
                result = (missing_status, "", "No target link found in the generated language site")
                work.append((row_index, str(key), lang, "", fallback_pages))
                checker.http_cache[f"__empty__:{row_index}:{lang}"] = result
            elif lang != "de" and link_type(str(source_href), site_origin) == "internal" and is_relative_internal_href(str(source_href)) and any(
                not is_relative_internal_href(href) or href != str(source_href) for href in hrefs
            ):
                result = ("relative link differs", "", f"Expected the German relative href unchanged: {source_href}")
                work.append((row_index, str(key), lang, "", fallback_pages))
                checker.http_cache[f"__empty__:{row_index}:{lang}"] = result
            else:
                for href in hrefs:
                    source_pages = pages_by_key[str(key)][lang].get(href)
                    if not source_pages:
                        source_pages = [
                            page
                            for page_list in pages_by_key[str(key)]["de"].values()
                            for page in page_list
                        ]
                    source_pages = list(dict.fromkeys(source_pages))
                    work.append((row_index, str(key), lang, href, source_pages))

    # Probe distinct external HTTP URLs by host. Internal URLs are resolved
    # against the generated release and never sent over the network.
    http_urls = set()
    for _, _, _, href, _ in work:
        if not href:
            continue
        parsed = urllib.parse.urlsplit(href)
        if parsed.scheme in {"http", "https"} and link_type(href, site_origin) == "external":
            http_urls.add(urllib.parse.urlunsplit(parsed._replace(fragment="")))
    urls_by_host: dict[str, list[str]] = defaultdict(list)
    for url in sorted(http_urls):
        host = urllib.parse.urlsplit(url).netloc.casefold()
        urls_by_host[host].append(url)

    def probe_host(urls: list[str]) -> dict[str, tuple[str, str, str]]:
        results = {}
        for index, url in enumerate(urls):
            result = checker.check_http(url)
            results[url] = result
            if result[0] == "HTTP 429":
                for deferred_url in urls[index + 1:]:
                    deferred = (
                        "HTTP 429 (host paused)",
                        "429",
                        f"Not requested: another URL on {urllib.parse.urlsplit(url).netloc} returned HTTP 429 after its delayed retry",
                    )
                    results[deferred_url] = deferred
                    checker.http_cache[deferred_url] = deferred
                break
        return results

    with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
        future_map = {executor.submit(probe_host, urls): host for host, urls in urls_by_host.items()}
        for future in as_completed(future_map):
            checker.http_cache.update(future.result())

    aggregates: dict[tuple[int, str], list[tuple[str, str, str]]] = defaultdict(list)
    details = []
    for row_index, key, lang, href, source_pages in work:
        if not href:
            result = checker.http_cache[f"__empty__:{row_index}:{lang}"]
            details.append((key, lang, "", *result))
        else:
            checked_urls = []
            for page in source_pages or ["index.html"]:
                checked_urls.append(checker.check(href, lang, [page]))
            # A URL shared across pages is valid only if each relative context resolves.
            states = {item[0] for item in checked_urls}
            if "not found" in states:
                result = ("not found", ";".join(dict.fromkeys(item[1] for item in checked_urls if item[1])), "; ".join(dict.fromkeys(item[2] for item in checked_urls if item[0] == "not found")))
            elif checked_urls and all(item[0] == "found" for item in checked_urls):
                result = ("found", ";".join(dict.fromkeys(item[1] for item in checked_urls if item[1])), "; ".join(dict.fromkeys(item[2] for item in checked_urls)))
            else:
                statuses = [item[0] for item in checked_urls if item[0] != "found"]
                status = statuses[0] if statuses else "local check inconclusive"
                result = (status, ";".join(dict.fromkeys(item[1] for item in checked_urls if item[1])), "; ".join(dict.fromkeys(item[2] for item in checked_urls if item[0] != "found")))
            details.append((key, lang, href, *result))
        aggregates[(row_index, lang)].append(result)

    summary: dict[str, int] = defaultdict(int)
    for row_index in range(2, links_sheet.max_row + 1):
        key = links_sheet.cell(row_index, link_headers["German link SHA-256"]).value
        if not key:
            continue
        for lang in LANGUAGES:
            prior = links_sheet.cell(row_index, link_headers[STATUS_COLUMNS[lang]]).value or ""
            results = aggregates.get((row_index, lang), [])
            states = {result[0] for result in results}
            if not results:
                status = "not found"
            elif "not found" in states:
                status = "not found"
            elif states.intersection({"missing link", "missing external link", "relative link differs"}):
                status = "; ".join(sorted(states.intersection({"missing link", "missing external link", "relative link differs"})))
            elif prior == "verified":
                status = "verified"
            elif states.issubset({"found"}):
                status = "found"
            else:
                status = "; ".join(sorted(states))
            links_sheet.cell(row_index, link_headers[STATUS_COLUMNS[lang]]).value = status
            summary[status] += 1

    if "Checks" in workbook.sheetnames:
        del workbook["Checks"]
    checks_sheet = workbook.create_sheet("Checks")
    checks_sheet.append(["German link SHA-256", "Language", "Link type", "Checked href", "Status", "HTTP response", "Detail", "Checked at UTC"])
    checked_at = datetime.now(UTC).isoformat(timespec="seconds")
    link_type_by_key = {
        str(links_sheet.cell(i, link_headers["German link SHA-256"]).value):
        link_type(str(links_sheet.cell(i, link_headers[LINK_COLUMNS["de"]]).value or ""), site_origin)
        for i in range(2, links_sheet.max_row + 1)
        if links_sheet.cell(i, link_headers["German link SHA-256"]).value
    }
    for row in details:
        key, lang, href, status, code, detail = row
        checks_sheet.append([key, lang, link_type_by_key.get(str(key), "unknown"), href, status, code, detail, checked_at])
    style_sheet(checks_sheet, [66, 14, 14, 70, 22, 16, 100, 28])
    workbook.save(workbook_path)
    workbook.close()
    return summary


def scan_command(args: argparse.Namespace) -> None:
    previous = args.previous_workbook
    if previous is None and args.output.exists():
        previous = args.output
    unique_count, occurrence_count = export_registry(args.site_root, args.output, previous)
    print(f"links={unique_count} occurrences={occurrence_count} workbook={args.output}")


def check_command(args: argparse.Namespace) -> None:
    summary = check_registry(args.workbook, args.site_root, args.site_origin, args.timeout, args.workers)
    print(" ".join(f"{key.replace(' ', '_')}={value}" for key, value in summary.items()))
    print(f"workbook={args.workbook}")


def external_report_command(args: argparse.Namespace) -> None:
    source = args.workbook.resolve()
    output = args.output.resolve()
    if source == output:
        raise ValueError("External report output must not overwrite the full link registry")
    source_workbook = load_workbook(source, read_only=True, data_only=True)
    required_sheets = {"Links", "Occurrences", "Checks", "Manifest"}
    if not required_sheets.issubset(source_workbook.sheetnames):
        raise ValueError("Run scan and check first; workbook must contain Links, Occurrences, Checks, and Manifest")

    source_links = source_workbook["Links"]
    link_header_values = next(source_links.iter_rows(values_only=True))
    link_headers = {value: index for index, value in enumerate(link_header_values) if value}
    if "Link type" not in link_headers or "German link SHA-256" not in link_headers:
        raise ValueError("Links sheet does not contain link type and German SHA-256 columns")
    external_keys = set()
    result_workbook = Workbook()
    result_workbook.remove(result_workbook.active)
    links = result_workbook.create_sheet("Links")
    links.append(list(link_header_values))
    for row in source_links.iter_rows(min_row=2, values_only=True):
        if row[link_headers["Link type"]] == "external":
            links.append(list(row))
            external_keys.add(row[link_headers["German link SHA-256"]])
    style_sheet(links, widths=[66, 50, 14, 40, 40, 40, 52, 52, 52, 52, 52, 22, 22, 22, 16])

    source_occurrences = source_workbook["Occurrences"]
    occurrence_header_values = next(source_occurrences.iter_rows(min_row=1, max_row=1, values_only=True))
    occurrence_headers = {value: index for index, value in enumerate(occurrence_header_values) if value}
    occurrences = result_workbook.create_sheet("Occurrences")
    occurrences.append(list(occurrence_header_values))
    occurrence_key_column = occurrence_headers["German link SHA-256"]
    for row in source_occurrences.iter_rows(min_row=2, values_only=True):
        if row[occurrence_key_column] in external_keys:
            occurrences.append(list(row))
    style_sheet(occurrences, widths=[66, 42, 16, 36, 52, 36, 52, 14, 36, 52, 14])

    source_checks = source_workbook["Checks"]
    checks_header_values = next(source_checks.iter_rows(min_row=1, max_row=1, values_only=True))
    checks_headers = {value: index for index, value in enumerate(checks_header_values) if value}
    if "Link type" not in checks_headers:
        raise ValueError("Checks sheet does not contain a Link type column; rerun check first")
    checks = result_workbook.create_sheet("Checks")
    checks.append(list(checks_header_values))
    checks_type_column = checks_headers["Link type"]
    for row in source_checks.iter_rows(min_row=2, values_only=True):
        if row[checks_type_column] == "external":
            checks.append(list(row))
    style_sheet(checks, widths=[66, 14, 14, 70, 26, 16, 100, 28])

    source_manifest = source_workbook["Manifest"]
    manifest = result_workbook.create_sheet("Manifest")
    for row in source_manifest.iter_rows(values_only=True):
        manifest.append(list(row))
    manifest.append(["Report scope", "External links only"])
    manifest.append(["External link keys", len(external_keys)])
    manifest.append(["External occurrences", occurrences.max_row - 1])
    style_sheet(manifest, widths=[32, 115])
    output.parent.mkdir(parents=True, exist_ok=True)
    result_workbook.save(output)
    source_workbook.close()
    result_workbook.close()
    print(f"external_link_keys={len(external_keys)} occurrences={occurrences.max_row - 1} report={output}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    scan = subparsers.add_parser("scan", help="Extract DE/FR/IT links into the Excel registry")
    scan.add_argument("--site-root", type=Path, default=DEFAULT_SITE_ROOT)
    scan.add_argument("--output", type=Path, default=DEFAULT_WORKBOOK)
    scan.add_argument("--previous-workbook", type=Path, help="Prior registry; verified mappings are reused by German SHA-256")
    scan.set_defaults(func=scan_command)

    check = subparsers.add_parser("check", help="Check registry links and update language statuses")
    check.add_argument("--site-root", type=Path, default=DEFAULT_SITE_ROOT)
    check.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    check.add_argument("--site-origin", default="https://50ohm.uska.ch", help="Origin for root-relative URLs outside language directories")
    check.add_argument("--timeout", type=float, default=12.0)
    check.add_argument("--workers", type=int, default=2, help="Concurrent external HTTP probes; low default avoids rate limiting")
    check.add_argument("--retry-429", type=int, default=1, help="Retries after HTTP 429, honoring Retry-After (default: 1)")
    check.add_argument("--max-retry-wait", type=float, default=30.0, help="Maximum seconds to wait before retrying HTTP 429")
    check.set_defaults(func=check_command)

    external = subparsers.add_parser("external-report", help="Create an Excel report containing external links only")
    external.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK, help="Full registry after scan and check")
    external.add_argument("--output", type=Path, default=DEFAULT_EXTERNAL_REPORT)
    external.set_defaults(func=external_report_command)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
