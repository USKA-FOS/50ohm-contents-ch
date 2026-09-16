#!/usr/bin/env python
"""Import the validated drawing-feedback workbook into localized TeX files.

The workbook is deliberately treated as a targeted patch. Existing localized
TeX files are preserved; a German source is used only when a target file does
not exist yet.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "work" / "drawing_feedback_review_2026-09-16.xlsx"
DEFAULT_AUDIT = ROOT / "work" / "drawing_feedback_review_2026-09-16.import-audit.json"
DEFAULT_RENDER_REPORT = ROOT / "work" / "drawing_text_audit" / "drawing_tex_translation_import_report.json"
SELECTED = "to_be_translated"
LANGUAGES = ("fr", "it")
REQUIRED = {
    "drawing", "canonical_id", "source_term_or_text", "proposal_fr",
    "proposal_it", "decision",
}


def rows_from_workbook(path: Path) -> list[dict[str, str]]:
    sheet = load_workbook(path, read_only=True, data_only=False)["Translation review"]
    values = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(values)]
    missing = sorted(REQUIRED - set(headers))
    if missing:
        raise ValueError(f"Workbook is missing columns: {missing}")
    return [
        {header: "" if value is None else str(value) for header, value in zip(headers, row)}
        for row in values
        if any(value is not None for value in row)
    ]


def normalize_tex(value: str) -> str:
    # Excel cells contain doubled backslashes when the review was prepared
    # from a Python representation of a TeX fragment.
    return (value or "").replace("\\\\", "\\").strip()


def update_meta(object_dir: Path, stem: str) -> None:
    path = object_dir / "object.meta.json"
    meta = json.loads(path.read_text(encoding="utf-8"))
    variants = meta.setdefault("language_variants", {})
    assets = meta.setdefault("metadata", {}).setdefault("language_asset", {})
    de_source = (assets.get("de.tex") or {}).get("source_path")
    de_source = de_source or (meta.get("metadata", {}).get("asset") or {}).get("tex_path")
    for language in LANGUAGES:
        variant = variants.setdefault(language, {})
        variant.setdefault("review_state", "to_be_reviewed")
        variant.setdefault("asset_files", {})["tex"] = f"{stem}.{language}.tex"
        assets[f"{language}.tex"] = {
            "canonical_file": f"{stem}.{language}.tex",
            "source_path": de_source,
        }
    path.write_text(json.dumps(meta, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def replace_all(text: str, source: str, target: str) -> tuple[str, int]:
    if not source:
        return text, 0
    count = text.count(source)
    return text.replace(source, target), count


def replace_abbreviation(text: str, source: str, target: str) -> tuple[str, int]:
    pattern = re.compile(rf"(?<![A-Z]){re.escape(source)}(?![A-Z])")
    return pattern.subn(target, text)


def apply_row(text: str, row: dict[str, str], language: str) -> tuple[str, int]:
    source = normalize_tex(row["source_term_or_text"])
    target = normalize_tex(row[f"proposal_{language}"])
    drawing = row["drawing"]

    if source == "Digitale Signalverarbeitung":
        # The German source deliberately splits this label over two nodes.
        source_parts = ("Digitale Signal-", "verarbeitung")
        target_parts = {
            "fr": ("traitement digital", "du signal"),
            "it": ("elaborazione digitale", "del segnale"),
        }[language]
        current_parts = {
            "fr": ("traitement numérique", "du signal"),
            "it": ("elaborazione digitale", "del segnale"),
        }[language]
        total = 0
        for old, new, current in zip(source_parts, target_parts, current_parts):
            text, count = replace_all(text, old, new)
            if not count:
                text, count = replace_all(text, current, new)
            total += count or int(new in text)
        return text, total

    if source == "N-Dotiert / P-Dotiert":
        source_parts = ("N-Dotiert", "P-Dotiert")
        target_parts = {
            "fr": ("Dopé N", "Dopé P"),
            "it": ("Drogato N", "Drogato P"),
        }[language]
        total = 0
        current_parts = {
            "fr": ("Dopé au N", "Dopé au P"),
            "it": ("Drogato con N", "Drogato con P"),
        }[language]
        for old, new, current in zip(source_parts, target_parts, current_parts):
            text, count = replace_all(text, old, new)
            if not count:
                text, count = replace_all(text, current, new)
            total += count or int(new in text)
        return text, total

    if source == "NF IN / NF OUT":
        # These labels are represented as NF_IN and NF_OUT in the TeX source.
        total = 0
        for old, new in ((r"\mathrm{NF}_\mathrm{IN}", r"\mathrm{BF}_\mathrm{IN}"),
                         (r"\mathrm{NF}_\mathrm{OUT}", r"\mathrm{BF}_\mathrm{OUT}")):
            text, count = replace_all(text, old, new)
            total += count or int(new in text)
        return text, total

    if drawing == "654" and source == "Numeral":
        current = {"fr": "numéral", "it": "numerale"}[language]
        text, count = replace_all(text, current, target)
        return text, count or int(target in text)

    if source == "Audioverstärker":
        rendered = {
            "fr": r"\mathrm{amplificateur}\ \mathrm{audio}",
            "it": r"\mathrm{amplificatore}\ \mathrm{audio}",
        }[language]
        old = r"\mathrm{Audioverstärker}"
        updated, count = replace_all(text, old, rendered)
        return updated, count or int(rendered in text)

    if source in {r"$f_\mathrm{T}$", r"$f_\text{T}$"}:
        # The source token is often followed by ``=...`` inside the same
        # math expression, so matching the complete ``$...$`` fragment is
        # too strict for entries such as ``$f_\mathrm{T}=...$``.
        target_inner = normalize_tex(target).strip("$")
        return re.subn(
            r"f_\\(?:mathrm|text)\{T\}",
            lambda _: target_inner,
            text,
        )

    if source in {"LW", "MW", "KW", "UKW"}:
        updated, count = replace_abbreviation(text, source, target)
        if count:
            return updated, count
        # The previous review used OL for the same German LW abbreviation.
        # Accept the corrected LF proposal without rebuilding the whole file.
        if source == "LW" and target == "LF":
            return replace_abbreviation(text, "OL", target)
        return text, int(target in text)

    # The reviewed workbook explicitly records this existing Italian wording.
    if drawing == "1044" and language == "it":
        text, count = replace_all(text, "Larghezza di banda", target)
        if count:
            return text, count

    return replace_all(text, source, target)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--audit", type=Path, default=DEFAULT_AUDIT)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument(
        "--decision-column",
        default="decision",
        help="Workbook column used to select rows (default: decision).",
    )
    args = parser.parse_args()

    if args.decision_column not in REQUIRED:
        REQUIRED.add(args.decision_column)
    rows = [row for row in rows_from_workbook(args.workbook)
            if row.get(args.decision_column, "").strip().lower() == SELECTED]
    if not rows:
        raise SystemExit(f"No rows with {args.decision_column}=to_be_translated")

    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        grouped[row["canonical_id"].strip()].append(row)

    audit = {
        "workflow": "drawing_feedback_review_import",
        "workbook": str(args.workbook),
        "decision_column": args.decision_column,
        "apply": args.apply,
        "selected_rows": len(rows),
        "drawings": [],
        "errors": [],
    }
    for canonical_id, drawing_rows in sorted(grouped.items()):
        object_dir = ROOT / "canonical" / "drawings" / canonical_id
        de_paths = sorted(object_dir.glob("*.de.tex"))
        if len(de_paths) != 1:
            audit["errors"].append(f"{canonical_id}: expected one .de.tex, found {len(de_paths)}")
            continue
        de_path = de_paths[0]
        stem = de_path.name.removesuffix(".de.tex")
        entry = {"canonical_id": canonical_id, "drawing": drawing_rows[0]["drawing"], "languages": {}}
        for language in LANGUAGES:
            target_path = object_dir / f"{stem}.{language}.tex"
            content = target_path.read_text(encoding="utf-8") if target_path.exists() else de_path.read_text(encoding="utf-8")
            original = content
            counts = []
            for row in drawing_rows:
                content, count = apply_row(content, row, language)
                counts.append({"source": row["source_term_or_text"], "count": count})
            missing = [
                item["source"]
                for item, row in zip(counts, drawing_rows)
                if item["count"] == 0
                and normalize_tex(row["proposal_" + language]) not in content
            ]
            entry["languages"][language] = {
                "path": str(target_path.relative_to(ROOT)),
                "existed": target_path.exists(),
                "changed": content != original,
                "replacement_counts": counts,
                "unmatched": missing,
            }
            if args.apply and content != original:
                target_path.write_text(content, encoding="utf-8")
        if args.apply and any(item["changed"] for item in entry["languages"].values()):
            update_meta(object_dir, stem)
        audit["drawings"].append(entry)

    args.audit.parent.mkdir(parents=True, exist_ok=True)
    args.audit.write_text(json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    if args.apply:
        render_report = {
            "workflow": "drawing_feedback_review_import",
            "updated_drawings": [
                {"canonical_reference": f"canonical/drawings/{entry['canonical_id']}"}
                for entry in audit["drawings"]
                if any(
                    (ROOT / language_info["path"]).exists()
                    for language_info in entry["languages"].values()
                )
            ],
        }
        DEFAULT_RENDER_REPORT.parent.mkdir(parents=True, exist_ok=True)
        DEFAULT_RENDER_REPORT.write_text(
            json.dumps(render_report, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
    print(f"selected_rows={len(rows)}")
    print(f"drawings={len(audit['drawings'])}")
    print(f"errors={len(audit['errors'])}")
    print(f"audit={args.audit}")
    if audit["errors"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
