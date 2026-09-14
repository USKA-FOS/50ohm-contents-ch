#!/usr/bin/env python
"""Import the reviewed fallback-drawing workbook into canonical TeX assets."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import extract_drawing_tex_translation_candidates as extractor
import import_drawing_tex_translations as tex_importer


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = REPO_ROOT / "work" / "drawing_text_audit" / "fallback_drawing_german_text_review_translation_proposals.xlsx"
DEFAULT_REVIEW_DIR = REPO_ROOT / "review" / "drawing_localization" / "2026-09-14"
DEFAULT_CSV = DEFAULT_REVIEW_DIR / "fallback_drawing_text_review.csv"
DEFAULT_AUDIT = DEFAULT_REVIEW_DIR / "fallback_drawing_text_import.audit.json"
DEFAULT_RENDER_REPORT = REPO_ROOT / "work" / "drawing_text_audit" / "drawing_tex_translation_import_report.json"
REQUIRED_COLUMNS = {
    "text_de_ori", "text", "fr", "it", "drawing_number", "canonical_id",
    "canonical_reference", "source_tex", "category", "decision",
}
SELECTED_DECISION = "to_be_translated"
LANGUAGES = ("fr", "it")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_review(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    sheet = load_workbook(path, read_only=True, data_only=False).active
    values = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value or "").strip() for value in next(values)]
    except StopIteration as error:
        raise ValueError(f"Review workbook is empty: {path}") from error
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    if missing:
        raise ValueError(f"Review workbook is missing columns: {missing}")
    rows = [
        {header: "" if value is None else str(value) for header, value in zip(headers, row)}
        for row in values
    ]
    return headers, rows


def write_csv(headers: list[str], rows: list[dict[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def build_candidate_index() -> dict[tuple[str, str], list[extractor.Candidate]]:
    result: dict[tuple[str, str], list[extractor.Candidate]] = defaultdict(list)
    for candidate in extractor.build_candidates():
        result[(candidate.canonical_reference, candidate.category)].append(candidate)
    return result


def matching_candidates(
    row: dict[str, str],
    index: dict[tuple[str, str], list[extractor.Candidate]],
) -> list[extractor.Candidate]:
    accepted = {
        tex_importer.normalize_lookup_term(row["text_de_ori"]),
        tex_importer.normalize_lookup_term(row["text"]),
    }
    return [
        candidate
        for candidate in index[(row["canonical_reference"].strip(), row["category"].strip())]
        if tex_importer.normalize_lookup_term(candidate.translatable_text) in accepted
    ]


def translated_fragment(candidate: extractor.Candidate, translation: str) -> str:
    normalized = tex_importer.normalize_translation_text(translation)
    segments = tex_importer.translation_segments_for_regular(candidate.raw_tex_fragment, normalized)
    protected = json.loads(candidate.protected_tokens)
    protected_by_segment = tex_importer.split_protected_tokens_by_segment(
        candidate.raw_tex_fragment, protected
    )
    return tex_importer.replace_fragment_text(
        candidate.raw_tex_fragment, segments, protected_tokens_by_segment=protected_by_segment
    )


def replace_once(text: str, source: str, target: str) -> tuple[str, bool]:
    if source not in text:
        return text, False
    return text.replace(source, target, 1), True


def replace_once_or_accept_applied(
    text: str, source: str, target: str
) -> tuple[str, bool, bool]:
    updated, replaced = replace_once(text, source, target)
    if replaced:
        return updated, True, False
    return text, False, target in text


def apply_candidate(text: str, candidate: extractor.Candidate, translation: str) -> tuple[str, bool]:
    category = candidate.category
    source_term = tex_importer.normalize_lookup_term(candidate.translatable_text)
    normalized = tex_importer.normalize_translation_text(translation)
    if category == "tikz_option_label":
        return tex_importer.replace_tikz_option_label(text, source_term, normalized)
    if category.startswith("pgfplots_") and category != "pgfplots_legend":
        option_name = category.removeprefix("pgfplots_")
        replacement = translated_fragment(candidate, translation)
        updated, replaced = tex_importer.replace_braced_option(
            text, option_name, candidate.raw_tex_fragment, replacement
        )
        if replaced:
            return updated, True
        return replace_once(text, candidate.raw_tex_fragment[1:-1], replacement[1:-1])
    if category == "pgfplots_legend":
        source = candidate.raw_tex_fragment[1:-1]
        if source_term not in source:
            return text, False
        target = source.replace(source_term, normalized, 1)
        return replace_once(text, source, target)
    if category == "circuitikz_bare_label":
        return tex_importer.replace_circuitikz_bare_label(text, source_term, normalized)
    if category == "math_text":
        return tex_importer.replace_math_text(text, source_term, normalized)
    return replace_once(text, candidate.raw_tex_fragment, translated_fragment(candidate, translation))


def validate_row_identity(row: dict[str, str]) -> tuple[Path, Path]:
    reference = row["canonical_reference"].strip()
    object_dir = REPO_ROOT / reference
    if not object_dir.is_dir() or object_dir.name != row["canonical_id"].strip():
        raise ValueError(f"Invalid canonical drawing reference: {reference}")
    source_path = object_dir / row["source_tex"].strip()
    if not source_path.is_file() or not source_path.name.endswith(".de.tex"):
        raise ValueError(f"Missing German TeX source: {source_path}")
    actual_number = extractor.load_figure_number(object_dir / "object.meta.json")
    if actual_number != row["drawing_number"].strip():
        raise ValueError(
            f"Drawing number mismatch for {reference}: workbook={row['drawing_number']} canonical={actual_number}"
        )
    return object_dir, source_path


def plan_import(rows: list[dict[str, str]]) -> dict[str, Any]:
    selected = [row for row in rows if row["decision"].strip().lower() == SELECTED_DECISION]
    if not selected:
        raise ValueError(f"No rows have decision={SELECTED_DECISION}")
    candidate_index = build_candidate_index()
    by_reference: dict[str, list[dict[str, str]]] = defaultdict(list)
    errors: list[str] = []
    seen: dict[tuple[str, str], tuple[str, str]] = {}
    for row in selected:
        reference = row["canonical_reference"].strip()
        identity = (reference, row["text_de_ori"].strip())
        translations = tuple(row[language].strip() for language in LANGUAGES)
        if not all(translations):
            errors.append(f"Missing translation for {identity}")
        if identity in seen and seen[identity] != translations:
            errors.append(f"Conflicting duplicate translations for {identity}")
        seen[identity] = translations
        try:
            validate_row_identity(row)
        except ValueError as error:
            errors.append(str(error))
        if row["category"] != "math_text_subscript" and not matching_candidates(row, candidate_index):
            errors.append(
                f"No current TeX candidate for {reference} / {row['category']} / {row['text_de_ori']}"
            )
        by_reference[reference].append(row)

    updated: list[dict[str, Any]] = []
    rendered_by_ref: dict[str, dict[str, str]] = {}
    for reference, reference_rows in sorted(by_reference.items()):
        object_dir, source_path = validate_row_identity(reference_rows[0])
        stem = source_path.name.removesuffix(".de.tex")
        rendered: dict[str, str] = {}
        replacement_counts = {language: 0 for language in LANGUAGES}
        already_applied_counts = {language: 0 for language in LANGUAGES}
        expected_counts = {language: 0 for language in LANGUAGES}
        for language in LANGUAGES:
            target_path = object_dir / f"{stem}.{language}.tex"
            osz_only = all(row["category"] == "math_text_subscript" for row in reference_rows)
            base_path = target_path if osz_only and target_path.is_file() else source_path
            content = base_path.read_text(encoding="utf-8")
            for row in reference_rows:
                translation = row[language].strip()
                if row["category"] == "math_text_subscript":
                    source = row["text_de_ori"].strip()
                    content, replaced, already_applied = replace_once_or_accept_applied(
                        content, source, tex_importer.normalize_translation_text(translation)
                    )
                    expected_counts[language] += 1
                    replacement_counts[language] += int(replaced)
                    already_applied_counts[language] += int(already_applied)
                    if not replaced and not already_applied:
                        errors.append(f"Missing OSZ expression in {base_path}: {source}")
                    continue
                matches = matching_candidates(row, candidate_index)
                expected_counts[language] += len(matches)
                for candidate in matches:
                    content, replaced = apply_candidate(content, candidate, translation)
                    replacement_counts[language] += int(replaced)
                    if not replaced:
                        errors.append(
                            f"Replacement failed in {base_path}: {row['category']} / {row['text_de_ori']}"
                        )
            rendered[language] = content
        rendered_by_ref[reference] = rendered
        updated.append(
            {
                "canonical_reference": reference,
                "drawing_number": reference_rows[0]["drawing_number"].strip(),
                "de_tex": str(source_path.relative_to(REPO_ROOT)),
                "fr_tex": str((object_dir / f"{stem}.fr.tex").relative_to(REPO_ROOT)),
                "it_tex": str((object_dir / f"{stem}.it.tex").relative_to(REPO_ROOT)),
                "selected_rows": len(reference_rows),
                "replacement_counts": replacement_counts,
                "already_applied_counts": already_applied_counts,
                "expected_counts": expected_counts,
            }
        )
    if errors:
        raise ValueError("Import validation failed:\n- " + "\n- ".join(errors))
    return {
        "selected_rows": len(selected),
        "updated_drawings": updated,
        "rendered_by_ref": rendered_by_ref,
    }


def apply_plan(plan: dict[str, Any]) -> None:
    for drawing in plan["updated_drawings"]:
        reference = drawing["canonical_reference"]
        object_dir = REPO_ROOT / reference
        stem = Path(drawing["de_tex"]).name.removesuffix(".de.tex")
        for language, content in plan["rendered_by_ref"][reference].items():
            (object_dir / f"{stem}.{language}.tex").write_text(content, encoding="utf-8")
        tex_importer.update_meta_for_tex(object_dir, stem)


def serializable_audit(
    plan: dict[str, Any], workbook: Path, csv_path: Path, applied: bool
) -> dict[str, Any]:
    return {
        "workflow": "fallback_drawing_text_review_import",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "applied": applied,
        "source_workbook": str(workbook.relative_to(REPO_ROOT)),
        "source_workbook_sha256": sha256(workbook),
        "review_csv": str(csv_path.relative_to(REPO_ROOT)),
        "review_csv_sha256": sha256(csv_path) if csv_path.exists() else None,
        "selected_rows": plan["selected_rows"],
        "updated_drawing_count": len(plan["updated_drawings"]),
        "line_break_markers": {
            "without_hyphen": "[[BR]]",
            "with_hyphen": "[[BR-]]",
        },
        "updated_drawings": plan["updated_drawings"],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--review-csv", type=Path, default=DEFAULT_CSV)
    parser.add_argument("--audit-json", type=Path, default=DEFAULT_AUDIT)
    parser.add_argument("--render-report", type=Path, default=DEFAULT_RENDER_REPORT)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    workbook = args.workbook.resolve()
    headers, rows = load_review(workbook)
    plan = plan_import(rows)
    print(f"worksheet_rows={len(rows)}")
    print(f"selected_rows={plan['selected_rows']}")
    print(f"updated_drawings={len(plan['updated_drawings'])}")
    if not args.apply:
        print("applied=false")
        return

    apply_plan(plan)
    write_csv(headers, rows, args.review_csv)
    audit = serializable_audit(plan, workbook, args.review_csv, applied=True)
    args.audit_json.parent.mkdir(parents=True, exist_ok=True)
    args.audit_json.write_text(json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    args.render_report.parent.mkdir(parents=True, exist_ok=True)
    args.render_report.write_text(
        json.dumps(
            {
                "workflow": audit["workflow"],
                "updated_drawings": plan["updated_drawings"],
                "languages": {language: len(plan["updated_drawings"]) for language in LANGUAGES},
            },
            indent=2,
            ensure_ascii=False,
        ) + "\n",
        encoding="utf-8",
    )
    print("applied=true")
    print(f"review_csv={args.review_csv}")
    print(f"audit={args.audit_json}")
    print(f"render_report={args.render_report}")


if __name__ == "__main__":
    main()
