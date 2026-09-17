"""Export German source changes as an operator review workbook.

This tool is intentionally read-only with regard to canonical/. It compares a
Git source revision with the current canonical model and creates the review
list consumed by the later import step.
"""

from __future__ import annotations

import argparse
import json
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from .import_incremental_german_source import (
        CANONICAL_ROOT,
        REPO_ROOT,
        TARGET_LANGUAGES,
        load_canonical_index,
        plan_structures,
        resolve_source_ref,
        scan_source,
        stable_id,
        unique_rename_matches,
    )
except ImportError:
    from import_incremental_german_source import (
        CANONICAL_ROOT,
        REPO_ROOT,
        TARGET_LANGUAGES,
        load_canonical_index,
        plan_structures,
        resolve_source_ref,
        scan_source,
        stable_id,
        unique_rename_matches,
    )


DEFAULT_OUTPUT = REPO_ROOT / "work" / "source_import_audits" / "german-source-import-review.xlsx"
HEADERS = (
    "object_id",
    "object_type",
    "source_key",
    "source_path",
    "previous_source_key",
    "change_kind",
    "changed_files",
    "source_commit",
    "source_commit_date",
    "translation_required",
    "tex_requires_explicit_decision",
    "decision",
)


def git_output(*args: str) -> str:
    return subprocess.check_output(["git", "-C", str(REPO_ROOT), *args], text=True).strip()


def source_commit_date(commit: str) -> str:
    return git_output("show", "-s", "--format=%cI", commit)


def stable_id_for_source(source_record: dict) -> str:
    family = source_record["family"]
    return stable_id(family.prefix, family.object_type, source_record["source_key"])


def build_rows(source_root: Path, source_commit: str) -> list[dict[str, object]]:
    source_records = scan_source(source_root)
    canonical_records = load_canonical_index(CANONICAL_ROOT)
    renames, _ = unique_rename_matches(source_records, canonical_records)
    matched_old = set(renames.values())
    rows: list[dict[str, object]] = []

    for identity, source_record in sorted(source_records.items()):
        old_identity = identity if identity in canonical_records else renames.get(identity)
        existing = canonical_records.get(old_identity) if old_identity else None
        if existing is None:
            change_kind = "new"
        elif old_identity != identity:
            change_kind = "updated"
        elif existing["fingerprint"] != source_record["fingerprint"]:
            change_kind = "updated"
        else:
            continue

        changed_suffixes = sorted(
            set(existing["files"]) | set(source_record["files"])
            if existing
            else set(source_record["files"])
        )
        if existing:
            changed_suffixes = [
                suffix
                for suffix in changed_suffixes
                if existing["files"].get(suffix) != source_record["files"].get(suffix)
            ]
        translation_required = any(suffix in {".md", ".html", ".txt", ".tex"} for suffix in changed_suffixes)
        has_tex = ".tex" in changed_suffixes
        html_change = ".html" in changed_suffixes
        decision = (
            ""
            if has_tex and change_kind != "new"
            else ("to_be_imported" if html_change or (change_kind == "new" and has_tex) else "")
        )
        rows.append(
            {
                "object_id": str(existing["meta"]["id"]) if existing else stable_id_for_source(source_record),
                "object_type": identity[0],
                "source_key": identity[1],
                "source_path": "; ".join(source_record["paths"].values()),
                "previous_source_key": old_identity[1] if old_identity and old_identity != identity else "",
                "change_kind": change_kind,
                "changed_files": ", ".join(changed_suffixes),
                "source_commit": source_commit,
                "source_commit_date": source_commit_date(source_commit),
                "translation_required": str(translation_required).lower(),
                "tex_requires_explicit_decision": str(has_tex and change_kind != "new").lower(),
                "decision": decision,
            }
        )

    for identity, record in sorted(canonical_records.items()):
        if identity in source_records or identity in matched_old:
            continue
        rows.append(
            {
                "object_id": str(record["meta"]["id"]),
                "object_type": identity[0],
                "source_key": identity[1],
                "source_path": str((record["meta"].get("source") or {}).get("path", "")),
                "previous_source_key": "",
                "change_kind": "deleted",
                "changed_files": ", ".join(sorted(record["files"])),
                "source_commit": source_commit,
                "source_commit_date": source_commit_date(source_commit),
                "translation_required": "false",
                "tex_requires_explicit_decision": str(".tex" in record["files"]).lower(),
                "decision": "",
            }
        )
    all_source_records = scan_source(source_root)
    all_canonical_records = load_canonical_index(CANONICAL_ROOT)
    all_renames, _ = unique_rename_matches(all_source_records, all_canonical_records)
    object_ids = {}
    for identity, source_record in all_source_records.items():
        old_identity = identity if identity in all_canonical_records else all_renames.get(identity)
        existing = all_canonical_records.get(old_identity) if old_identity else None
        object_ids[identity] = str(existing["meta"]["id"]) if existing else stable_id_for_source(source_record)
    structure_entries, structure_payloads, _, _ = plan_structures(source_root, CANONICAL_ROOT, object_ids)
    for entry in structure_entries:
        if entry["action"] == "unchanged":
            continue
        edition = str(entry["edition"])
        payload = structure_payloads.get(edition) or {}
        meta = payload.get("meta") or {}
        change_kind = "new" if entry["action"] == "added" else "updated"
        rows.append(
            {
                "object_id": str(meta.get("id", f"edition_{edition}")),
                "object_type": "curriculum_structure",
                "source_key": edition,
                "source_path": f"toc/{edition}.json",
                "previous_source_key": "",
                "change_kind": change_kind,
                "changed_files": ".json",
                "source_commit": source_commit,
                "source_commit_date": source_commit_date(source_commit),
                "translation_required": "true",
                "tex_requires_explicit_decision": "false",
                "decision": "to_be_imported" if change_kind == "new" else "",
            }
        )
    return rows


def write_workbook(rows: list[dict[str, object]], output: Path, source_ref: str, source_commit: str) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "source_review"
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append([row[header] for header in HEADERS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [16, 20, 34, 64, 24, 14, 18, 42, 24, 18, 28, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    metadata = workbook.create_sheet("metadata")
    metadata.append(["key", "value"])
    metadata.append(["source_ref", source_ref])
    metadata.append(["source_commit", source_commit])
    metadata.append(["canonical_authority", "canonical/"])
    metadata.append(["import_policy", "Review and approve this workbook before any canonical import."])
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export German source changes for explicit review.")
    parser.add_argument("--source-ref", default="origin/review/de/main")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    source_root, commit, temporary = resolve_source_ref(args.source_ref)
    try:
        rows = build_rows(source_root, commit)
        write_workbook(rows, args.output.resolve(), args.source_ref, commit)
    finally:
        temporary.cleanup()
    counts: dict[str, int] = {}
    for row in rows:
        kind = str(row["change_kind"])
        counts[kind] = counts.get(kind, 0) + 1
    print(f"workbook={args.output.resolve()}")
    print(f"source_ref={args.source_ref}")
    print(f"source_commit={commit}")
    print(json.dumps(counts, sort_keys=True))
    print(f"candidate_count={len(rows)}")


if __name__ == "__main__":
    main()
