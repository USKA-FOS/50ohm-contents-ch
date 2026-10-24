from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
import pytest

from tools.prepare_quotation_manual_review import excerpt, prepare


def test_prepare_preserves_exact_copies_and_reports_each_quote_with_its_line(tmp_path):
    canonical = tmp_path / "canonical"
    folder = canonical / "sections" / "sc_sample"
    folder.mkdir(parents=True)
    texts = {
        "de": '% "Not a public quotation"\n"Epsilon R"\n\n„vollständig“\n',
        "fr": '«\u00a0epsilon R\u00a0»\n\ncomplètement\n',
        "it": '«Epsilon R»\n\n«completamente»\n',
    }
    for language, text in texts.items():
        (folder / f"body.{language}.md").write_bytes(text.encode())
    (folder / "object.meta.json").write_text(json.dumps({"id": "sc_sample", "source": {"path": "contents/sections/kondensator.md"}}))
    candidates = tmp_path / "candidates.json"
    candidates.write_text(json.dumps({"post_change_quote_pair_count_screen": {
        "candidate_count": 1, "candidates": [{"object": "sections/sc_sample", "paired_quote_count": {"de": 2, "fr": 1, "it": 2}}],
    }}))
    output = tmp_path / "review"
    result = prepare(canonical_root=canonical, candidates_manifest=candidates, output_dir=output,
                     translation_tool_root=Path(__file__).resolve().parents[2] / "50ohm-ai-translation-tool")
    assert result["object_count"] == 1 and result["quotation_rows"] == 2
    workbook = load_workbook(output / "quotation-review.xlsx")
    sheet = workbook["Guillemets"]
    assert sheet.max_column == 5
    assert [sheet.cell(2, c).value for c in (2, 3, 4, 5)] == ["kondensator.md", "Epsilon R", "epsilon R", "Epsilon R"]
    assert [sheet.cell(3, c).value for c in (3, 4, 5)] == ["vollständig", "—", "completamente"]
    assert "ligne 4" in sheet["C3"].comment.text
    assert sheet["C3"].hyperlink.target == "sc_sample/body.de.md"
    for language, text in texts.items():
        entry = result["objects"][0]["files"][language]
        data = text.encode()
        assert (folder / f"body.{language}.md").read_bytes() == data
        assert (output / entry["review_path"]).read_bytes() == data
        assert entry["sha256"] == hashlib.sha256(data).hexdigest()
        assert entry["original_text"] == text
    edited = output / "sc_sample/body.fr.md"
    edited.write_text("Human correction")
    with pytest.raises(ValueError, match="must be new"):
        prepare(canonical_root=canonical, candidates_manifest=candidates, output_dir=output,
                translation_tool_root=Path(__file__).resolve().parents[2] / "50ohm-ai-translation-tool")
    assert edited.read_text() == "Human correction"


def test_excerpt_limits_visible_text_to_fifteen_characters():
    text = "é" * 16
    assert excerpt(text) == "é" * 15 + "…"
    assert excerpt("A short phrase") == "A short phrase"
