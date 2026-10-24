from email.message import Message
from pathlib import Path
import urllib.error
from argparse import Namespace

from openpyxl import Workbook, load_workbook

from tools.manage_site_link_registry import (
    Anchor,
    AnchorParser,
    LinkChecker,
    align_anchors,
    ascii_http_url,
    external_report_command,
    is_reveal_route_fragment,
    is_relative_internal_href,
    link_type,
    link_hash,
)
import tools.manage_site_link_registry as registry


def test_anchor_parser_reads_nested_text_entities_and_image_alt():
    parser = AnchorParser()
    parser.feed('<a href="https://example.test/a"><strong>Pr&uuml;fung</strong> <img alt="PDF"></a>')

    assert len(parser.anchors) == 1
    assert parser.anchors[0].href == "https://example.test/a"
    assert parser.anchors[0].text == "Prüfung PDF"


def test_hash_uses_german_href_and_scopes_bare_fragments_to_page():
    assert link_hash("https://example.test/a", "one.html") == link_hash("https://example.test/a", "two.html")
    assert link_hash("#part", "one.html") != link_hash("#part", "two.html")
    assert link_hash("page.html#part", "one.html") == link_hash("page.html#part", "two.html")


def test_internal_and_external_links_are_classified_separately():
    origin = "https://50ohm.uska.ch"
    assert link_type("../fr/index.html", origin) == "internal"
    assert link_type("/fr/index.html", origin) == "internal"
    assert link_type("https://50ohm.uska.ch/isq/site/de/index.html", origin) == "internal"
    assert link_type("https://example.org/info", origin) == "external"
    assert is_relative_internal_href("../fr/index.html")
    assert is_relative_internal_href("#chapter")
    assert not is_relative_internal_href("https://example.org/info")


def test_same_site_absolute_link_is_checked_locally_without_http(tmp_path, monkeypatch):
    (tmp_path / "de").mkdir()
    (tmp_path / "de" / "page.html").write_text('<h1 id="part">Part</h1>', encoding="utf-8")
    checker = LinkChecker(tmp_path, "https://50ohm.uska.ch", timeout=1)
    monkeypatch.setattr(checker, "check_http", lambda url: (_ for _ in ()).throw(AssertionError(url)))

    assert checker.check("https://50ohm.uska.ch/isq/site/de/page.html#part", "de", ["index.html"])[0] == "found"


def test_http_429_waits_then_retries_once(monkeypatch):
    calls = []
    waits = []

    class Response:
        status = 200

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self, _size):
            return b"x"

        def geturl(self):
            return "https://example.org/"

    def fake_urlopen(_request, timeout):
        calls.append(timeout)
        if len(calls) == 1:
            headers = Message()
            headers["Retry-After"] = "3"
            raise urllib.error.HTTPError("https://example.org/", 429, "Too Many Requests", headers, None)
        return Response()

    monkeypatch.setattr(registry.urllib.request, "urlopen", fake_urlopen)
    monkeypatch.setattr(registry.time, "sleep", waits.append)
    checker = LinkChecker(Path("."), "https://50ohm.uska.ch", timeout=5, retry_429=1, max_retry_wait=30)

    assert checker.check_http("https://example.org/")[0:2] == ("found", "200")
    assert calls == [5, 5]
    assert waits == [3.0]


def test_http_429_is_not_retried_before_a_long_retry_after(monkeypatch):
    calls = []

    def fake_urlopen(_request, timeout):
        calls.append(timeout)
        headers = Message()
        headers["Retry-After"] = "120"
        raise urllib.error.HTTPError("https://example.org/", 429, "Too Many Requests", headers, None)

    monkeypatch.setattr(registry.urllib.request, "urlopen", fake_urlopen)
    checker = LinkChecker(Path("."), "https://50ohm.uska.ch", timeout=5, retry_429=1, max_retry_wait=30)

    assert checker.check_http("https://example.org/")[0:2] == ("HTTP 429", "429")
    assert calls == [5]


def test_external_report_contains_no_internal_rows(tmp_path):
    source = tmp_path / "full.xlsx"
    output = tmp_path / "external.xlsx"
    workbook = Workbook()
    links = workbook.active
    links.title = "Links"
    links.append(["German link SHA-256", "Link type", "Link DE", "Status DE"])
    links.append(["internal-key", "internal", "../de/page.html", "found"])
    links.append(["external-key", "external", "https://example.org/", "HTTP 429"])
    occurrences = workbook.create_sheet("Occurrences")
    occurrences.append(["German link SHA-256", "German page"])
    occurrences.append(["internal-key", "de/index.html"])
    occurrences.append(["external-key", "de/index.html"])
    checks = workbook.create_sheet("Checks")
    checks.append(["German link SHA-256", "Link type", "Status"])
    checks.append(["internal-key", "internal", "found"])
    checks.append(["external-key", "external", "HTTP 429"])
    manifest = workbook.create_sheet("Manifest")
    manifest.append(["Field", "Value"])
    workbook.save(source)

    external_report_command(Namespace(workbook=source, output=output))
    result = load_workbook(output, read_only=True, data_only=True)
    assert [row[0] for row in result["Links"].iter_rows(min_row=2, values_only=True)] == ["external-key"]
    assert [row[0] for row in result["Occurrences"].iter_rows(min_row=2, values_only=True)] == ["external-key"]
    assert [row[0] for row in result["Checks"].iter_rows(min_row=2, values_only=True)] == ["external-key"]


def test_http_url_encodes_unicode_path_and_international_hostname():
    assert ascii_http_url("https://münich.example/prüfen?q=français") == (
        "https://xn--mnich-kva.example/pr%C3%BCfen?q=fran%C3%A7ais"
    )


def test_reveal_slide_route_fragments_are_not_html_ids():
    assert is_reveal_route_fragment("13")
    assert is_reveal_route_fragment("13/2")
    assert not is_reveal_route_fragment("ref_a_psk")


def test_alignment_does_not_shift_links_after_missing_target_anchors():
    source = [Anchor(f"target-{index}") for index in range(8)]
    target = [source[0], source[1], source[2], source[6], source[7]]

    assert align_anchors(source, target) == {0: 0, 1: 1, 2: 2, 6: 3, 7: 4}
